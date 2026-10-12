import re

def test_fields_endpoint(client):
    response = client.get("/fields/")

    assert response.status_code == 200

    fields = response.json()["fields"]

    assert "province" in fields
    assert "city" in fields
    assert "car_company" in fields
    assert "car_name" in fields
    assert "mobile_number" in fields
    assert "postal_code" in fields


def test_generate_location(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "locations",
            "fields": [
                {
                    "type": "province",
                    "name": "province",
                },
                {
                    "type": "city",
                    "name": "city",
                },
            ],
            "count": 5,
            "format": "json",
        },
    )

    assert response.status_code == 200

    rows = response.json()["data"]

    assert len(rows) == 5

    for row in rows:
        assert row["province"] == "تهران"
        assert row["city"] == "کرج تست"


def test_generate_car(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "vehicles",
            "fields": [
                {
                    "type": "car_company",
                    "name": "manufacturer",
                },
                {
                    "type": "car_name",
                    "name": "model",
                },
            ],
            "count": 5,
            "format": "json",
        },
    )

    assert response.status_code == 200

    for row in response.json()["data"]:
        assert row["manufacturer"] == "ایران‌خودرو"
        assert row["model"] == "دنا"


def test_unique_email_username(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "users",
            "fields": [
                {
                    "type": "email",
                    "name": "email",
                },
                {
                    "type": "username",
                    "name": "username",
                },
            ],
            "count": 30,
            "format": "json",
        },
    )

    assert response.status_code == 200

    rows = response.json()["data"]

    emails = [r["email"] for r in rows]
    usernames = [r["username"] for r in rows]

    assert len(set(emails)) == 30
    assert len(set(usernames)) == 30

    for email in emails:
        assert re.fullmatch(
            r"ali[0-9]*@example\.com",
            email,
        )


def test_generate_sql(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "vehicles",
            "fields": [
                {
                    "type": "car_name",
                    "name": "model",
                },
                {
                    "type": "car_year",
                    "name": "model_year",
                },
            ],
            "count": 2,
            "format": "sql",
            "include_create_table": True,
        },
    )

    assert response.status_code == 200

    content = response.text

    assert "CREATE TABLE IF NOT EXISTS" in content
    assert 'INSERT INTO "vehicles"' in content
    assert '"model_year" INTEGER' in content
    assert '"id" INTEGER PRIMARY KEY' in content


def test_generate_csv(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "mobiles",
            "fields": [
                {
                    "type": "mobile_number",
                    "name": "mobile",
                },
            ],
            "count": 3,
            "format": "csv",
        },
    )

    assert response.status_code == 200

    content = response.content.decode("utf-8-sig")

    lines = content.splitlines()

    assert lines[0] == "id,mobile"
    assert len(lines) == 4

    for line in lines[1:]:
        _, mobile = line.split(",")
        assert re.fullmatch(r"09[0-9]{9}", mobile)


def test_invalid_api_request(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "users",
            "fields": [
                {
                    "type": "invalid",
                    "name": "bad",
                }
            ],
            "count": 5,
            "format": "json",
        },
    )

    assert response.status_code == 422


def test_generate_excel(client):
    response = client.post(
        "/generate/",
        json={
            "table_name": "postal_data",
            "fields": [
                {
                    "type": "postal_code",
                    "name": "postal_code",
                },
            ],
            "count": 3,
            "format": "excel",
        },
    )

    assert response.status_code == 200

    assert response.content[:2] == b"PK"

    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(response.content)
    )

    sheet = workbook.active

    assert sheet["B1"].value == "postal_code"

    for row in range(2, 5):
        value = sheet.cell(row, 2).value

        assert isinstance(value, str)
        assert re.fullmatch(r"[0-9]{10}", value)

    workbook.close()
